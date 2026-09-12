"""Column-level Lineage IR from the formula pass."""

from __future__ import annotations

from collections.abc import Mapping

from openpyxl.utils.cell import get_column_letter

from exmigrate.analyzer.formulas import SheetFormulas, SheetRef
from exmigrate.analyzer.keys import table_by_sheet
from exmigrate.contracts.ir import SchemaIR
from exmigrate.contracts.lineage import LineageEdge, LineageIR, LineageNode

OP_DYNAMIC = "dynamic"
DYNAMIC_LABEL = "INDIRECT / OFFSET (unresolved)"


def build_lineage(ir: SchemaIR, facts: Mapping[str, Mapping[str, SheetFormulas]]) -> LineageIR:
    """Column-level lineage for ``facts`` (file name → sheet title → formula facts)."""
    nodes: dict[str, LineageNode] = {}
    edges: dict[tuple[str, str, str], int] = {}
    files = {t.source_file for t in ir.tables}

    for file, sheets in facts.items():
        for sheet, sheet_facts in sheets.items():
            table = table_by_sheet(ir, file, sheet)
            if table is None:
                continue
            for col_idx, flows in sheet_facts.flows.items():
                if col_idx >= len(table.columns):
                    continue
                dst = _column_node(ir, SheetRef(sheet, col_idx), file, files)
                nodes.setdefault(dst.id, dst)
                for (ref, op), count in flows.items():
                    src = _column_node(ir, ref, file, files)
                    nodes.setdefault(src.id, src)
                    key = (src.id, dst.id, op)
                    edges[key] = edges.get(key, 0) + count
            for col_idx, count in sheet_facts.dynamic.items():
                if col_idx >= len(table.columns):
                    continue
                dst = _column_node(ir, SheetRef(sheet, col_idx), file, files)
                nodes.setdefault(dst.id, dst)
                dyn = LineageNode(
                    id=f"dynamic:{file}", kind="dynamic_reference", label=DYNAMIC_LABEL, file=file
                )
                nodes.setdefault(dyn.id, dyn)
                key = (dyn.id, dst.id, OP_DYNAMIC)
                edges[key] = edges.get(key, 0) + count

    return LineageIR(
        level="column",
        nodes=list(nodes.values()),
        edges=[
            LineageEdge(source=s, target=t, op=op, formula_count=n)
            for (s, t, op), n in edges.items()
        ],
    )


def _column_node(ir: SchemaIR, ref: SheetRef, current_file: str, files: set[str]) -> LineageNode:
    file = ref.book or current_file
    table = table_by_sheet(ir, file, ref.sheet)
    missing = file not in files
    if table is not None and ref.column < len(table.columns):
        column = table.columns[ref.column].name
        sheet_label = table.name
    else:
        column = get_column_letter(ref.column + 1)
        sheet_label = ref.sheet
    return LineageNode(
        id=f"column:{file}/{ref.sheet}/{column}",
        kind="column",
        label=f"{sheet_label}.{column}" + (" (missing)" if missing else ""),
        file=file,
        sheet=ref.sheet,
        column=column,
        missing=missing,
    )
