"""Second read pass (``data_only=False``): formula strings.

Recognised formula families: plain references (``=A2``, ``=Sheet2!B3``,
``=[1]Other!C4``), arithmetic over them, lookups (``VLOOKUP`` / ``XLOOKUP``)
and conditional aggregates (``SUMIF(S)`` / ``COUNTIF(S)`` / ``AVERAGEIF(S)``)
whose range lives in another sheet or in an external workbook. ``INDIRECT`` /
``OFFSET`` cannot be resolved statically and are only flagged as dynamic.
Anything else is treated as an opaque formula: the column still counts as
derived, but yields no key evidence.
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path, PurePosixPath
from urllib.parse import unquote

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Token, Tokenizer, TokenizerError
from openpyxl.utils.cell import column_index_from_string
from openpyxl.workbook.external_link.external import ExternalLink
from openpyxl.workbook.workbook import Workbook

DERIVED_MIN_RATIO = 0.5

JOIN_FUNCS = frozenset({"VLOOKUP", "XLOOKUP", "HLOOKUP", "LOOKUP", "INDEX", "MATCH", "XMATCH"})
AGGREGATE_FUNCS = frozenset(
    {"SUMIF", "SUMIFS", "COUNTIF", "COUNTIFS", "AVERAGEIF", "AVERAGEIFS", "MINIFS", "MAXIFS"}
)
DYNAMIC_FUNCS = frozenset({"INDIRECT", "OFFSET"})

OP_JOIN = "join"
OP_AGGREGATE = "aggregate"
OP_CALCULATE = "calculate"
OP_COPY = "copy"

_CELL_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+)$")
_RANGE_RE = re.compile(r"^\$?([A-Za-z]{1,3})(?:\$?\d+)?(?::\$?([A-Za-z]{1,3})(?:\$?\d+)?)?$")
_BOOK_RE = re.compile(r"^\[(\d+)\](.*)$")
_FUNC_RE = re.compile(
    r"(?<![A-Za-z0-9_])(VLOOKUP|XLOOKUP|SUMIFS?|COUNTIFS?|AVERAGEIFS?)\s*\(", re.IGNORECASE
)
_ANY_FUNC_RE = re.compile(r"(?<![A-Za-z0-9_.!])(?:_xl[a-z]+\.)?([A-Za-z][A-Za-z0-9_.]*)\s*\(")


@dataclass(frozen=True)
class SheetRef:
    """A (sheet, zero-based column) location; ``book`` names an external workbook."""

    sheet: str
    column: int
    book: str | None = None

    @property
    def sheet_label(self) -> str:
        """``sheet`` or ``[book]sheet`` for external references."""
        return f"[{self.book}]{self.sheet}" if self.book else self.sheet


@dataclass(frozen=True)
class LookupEdge:
    """``lookup_column`` (in ``sheet``) is looked up against ``target``."""

    sheet: str
    lookup_column: int
    target: SheetRef
    function: str


@dataclass
class SheetFormulas:
    """Formula facts for one worksheet, keyed by zero-based column."""

    title: str
    formula_counts: dict[int, int] = field(default_factory=dict)
    samples: dict[int, str] = field(default_factory=dict)
    lookups: list[LookupEdge] = field(default_factory=list)
    references: dict[int, set[SheetRef]] = field(default_factory=dict)
    functions: dict[int, dict[str, int]] = field(default_factory=dict)
    flows: dict[int, dict[tuple[SheetRef, str], int]] = field(default_factory=dict)
    dynamic: dict[int, int] = field(default_factory=dict)

    def derived_columns(self, row_count: int) -> set[int]:
        """Columns where at least ``DERIVED_MIN_RATIO`` of rows hold a formula."""
        if row_count <= 0:
            return set()
        return {
            col
            for col, count in self.formula_counts.items()
            if count / row_count >= DERIVED_MIN_RATIO
        }


def scan_formulas(path: str | Path, header_rows: dict[str, int]) -> dict[str, SheetFormulas]:
    """Collect formula facts per sheet (see :func:`scan_workbook`)."""
    return scan_workbook(path, header_rows)[0]


def scan_workbook(
    path: str | Path, header_rows: dict[str, int]
) -> tuple[dict[str, SheetFormulas], list[str]]:
    """Collect formula facts per sheet plus the workbook's external link names.

    ``header_rows`` maps sheet title → 1-based header row; rows at or above the
    header are ignored. Sheets missing from the mapping are skipped.
    """
    wb = load_workbook(path, data_only=False, read_only=True)
    out: dict[str, SheetFormulas] = {}
    try:
        books = external_books(wb)
        for ws in wb.worksheets:
            header = header_rows.get(ws.title)
            if header is None:
                continue
            facts = SheetFormulas(title=ws.title)
            for row in ws.iter_rows(min_row=header + 1, values_only=True):
                for col, value in enumerate(row):
                    if not isinstance(value, str) or not value.startswith("="):
                        continue
                    facts.formula_counts[col] = facts.formula_counts.get(col, 0) + 1
                    names = function_names(value)
                    counts = facts.functions.setdefault(col, {})
                    for name in names:
                        counts[name] = counts.get(name, 0) + 1
                    facts.samples.setdefault(col, value)
                    _extract(value, names, ws.title, col, facts, books)
            out[ws.title] = facts
    finally:
        wb.close()
    return out, books


def external_books(wb: Workbook) -> list[str]:
    """File names of the workbook's external links, in ``[1]``, ``[2]``… order."""
    names: list[str] = []
    # openpyxl keeps parsed xl/externalLinks/*.xml here; there is no public accessor.
    links: list[ExternalLink] = wb._external_links  # type: ignore[attr-defined]
    for link in links:
        target = link.file_link.Target if link.file_link is not None else ""
        names.append(unquote(PurePosixPath(str(target).replace("\\", "/")).name))
    return names


def function_names(formula: str) -> list[str]:
    """Upper-cased function names called in ``formula`` (``_xlfn.`` prefixes stripped)."""
    stripped = re.sub(r"\"(?:[^\"]|\"\")*\"|'(?:[^']|'')*'", "''", formula)
    return [m.group(1).upper() for m in _ANY_FUNC_RE.finditer(stripped)]


def classify(formula: str, names: Sequence[str]) -> str:
    """Edge semantics of ``formula``: join, aggregate, copy or calculate."""
    upper = set(names)
    if upper & JOIN_FUNCS:
        return OP_JOIN
    if upper & AGGREGATE_FUNCS:
        return OP_AGGREGATE
    tokens = _tokens(formula)
    if tokens is not None and len(tokens) == 1 and _is_range(tokens[0]):
        return OP_COPY
    if tokens is None and _REF_TOKEN_RE.fullmatch(formula[1:].strip()):
        return OP_COPY
    return OP_CALCULATE


def _extract(
    formula: str,
    names: Sequence[str],
    sheet: str,
    col: int,
    facts: SheetFormulas,
    books: Sequence[str],
) -> None:
    refs = _sheet_refs(formula, sheet, books)
    for match in _FUNC_RE.finditer(formula):
        args = _split_args(formula, match.end())
        func = match.group(1).upper()
        edges = (
            _aggregate_edges(func, args, sheet, books)
            if func in AGGREGATE_FUNCS
            else [_lookup_edge(func, args, sheet, books)]
        )
        for edge in edges:
            if edge is not None and edge not in facts.lookups:
                facts.lookups.append(edge)
        returned = _returned_column(func, args, sheet, books)
        if returned is not None:
            refs.add(returned)
    foreign = {r for r in refs if r.book is not None or r.sheet != sheet}
    if foreign:
        facts.references.setdefault(col, set()).update(foreign)
    if refs:
        op = classify(formula, names)
        flows = facts.flows.setdefault(col, {})
        for ref in refs:
            flows[(ref, op)] = flows.get((ref, op), 0) + 1
    if set(names) & DYNAMIC_FUNCS:
        facts.dynamic[col] = facts.dynamic.get(col, 0) + 1


def _lookup_edge(
    func: str, args: Sequence[str], sheet: str, books: Sequence[str] = ()
) -> LookupEdge | None:
    if len(args) < 2:
        return None
    return _edge(func, args[0], args[1], sheet, books)


def _returned_column(
    func: str, args: Sequence[str], sheet: str, books: Sequence[str]
) -> SheetRef | None:
    """Column a VLOOKUP returns when its index argument is a literal (HLOOKUP returns rows)."""
    if func != "VLOOKUP" or len(args) < 3 or not args[2].isdigit():
        return None
    first = _range_first_column(args[1], sheet, books)
    if first is None:
        return None
    return SheetRef(first.sheet, first.column + int(args[2]) - 1, first.book)


def _aggregate_edges(
    func: str, args: Sequence[str], sheet: str, books: Sequence[str]
) -> list[LookupEdge | None]:
    """Criteria cell → criteria range pairs of a conditional aggregate."""
    pairs: list[tuple[str, str]]
    if func in {"SUMIF", "COUNTIF", "AVERAGEIF"}:
        pairs = [(args[1], args[0])] if len(args) >= 2 else []
    else:
        start = 0 if func == "COUNTIFS" else 1
        pairs = [
            (args[i + 1], args[i]) for i in range(start, len(args) - 1, 2) if i + 1 < len(args)
        ]
    return [_edge(func, criteria, rng, sheet, books) for criteria, rng in pairs]


def _edge(
    func: str, lookup_arg: str, array_arg: str, sheet: str, books: Sequence[str]
) -> LookupEdge | None:
    lookup = _cell_column(lookup_arg, sheet, books)
    if lookup is None or lookup.book is not None or lookup.sheet != sheet:
        return None
    target = _range_first_column(array_arg, sheet, books)
    if target is None or (target.book is None and target.sheet == sheet):
        return None
    return LookupEdge(sheet=sheet, lookup_column=lookup.column, target=target, function=func)


def _split_args(formula: str, start: int) -> list[str]:
    """Split the argument list of a call whose ``(`` ends at ``start``."""
    args: list[str] = []
    depth = 0
    in_str = False
    buf: list[str] = []
    for ch in formula[start:]:
        if in_str:
            buf.append(ch)
            if ch == '"':
                in_str = False
            continue
        if ch == '"':
            in_str = True
            buf.append(ch)
        elif ch == "(":
            depth += 1
            buf.append(ch)
        elif ch == ")":
            if depth == 0:
                args.append("".join(buf).strip())
                return args
            depth -= 1
            buf.append(ch)
        elif ch in ",;" and depth == 0:
            args.append("".join(buf).strip())
            buf = []
        else:
            buf.append(ch)
    args.append("".join(buf).strip())
    return args


def _split_sheet(
    ref: str, default_sheet: str, books: Sequence[str]
) -> tuple[str | None, str, str] | None:
    """Split ``[1]Sheet!A1`` into (book, sheet, cell); ``None`` for unknown books."""
    if "!" not in ref:
        return None, default_sheet, ref
    sheet, _, rest = ref.rpartition("!")
    sheet = sheet.strip()
    if sheet.startswith("'") and sheet.endswith("'"):
        sheet = sheet[1:-1].replace("''", "'")
    book: str | None = None
    m = _BOOK_RE.match(sheet)
    if m:
        index = int(m.group(1))
        if not 1 <= index <= len(books):
            return None
        book, sheet = books[index - 1], m.group(2)
    return book, sheet, rest.strip()


def _cell_column(ref: str, default_sheet: str, books: Sequence[str] = ()) -> SheetRef | None:
    parts = _split_sheet(ref.strip(), default_sheet, books)
    if parts is None:
        return None
    book, sheet, rest = parts
    m = _CELL_RE.match(rest)
    if not m:
        return None
    return SheetRef(sheet=sheet, column=column_index_from_string(m.group(1).upper()) - 1, book=book)


def _range_first_column(ref: str, default_sheet: str, books: Sequence[str] = ()) -> SheetRef | None:
    parts = _split_sheet(ref.strip(), default_sheet, books)
    if parts is None:
        return None
    book, sheet, rest = parts
    m = _RANGE_RE.match(rest)
    if not m:
        return None
    return SheetRef(sheet=sheet, column=column_index_from_string(m.group(1).upper()) - 1, book=book)


_REF_TOKEN_RE = re.compile(
    r"(?:(?:'(?:[^']|'')+'|[A-Za-z0-9_.\[\]]+)!)?\$?[A-Za-z]{1,3}\$?\d*(?::\$?[A-Za-z]{1,3}\$?\d*)?"
)


def _tokens(formula: str) -> list[Token] | None:
    try:
        return [t for t in Tokenizer(formula).items if t.type != Token.WSPACE]
    except (TokenizerError, IndexError):
        return None


def _is_range(token: Token) -> bool:
    return token.type == Token.OPERAND and token.subtype == Token.RANGE


def _sheet_refs(formula: str, default_sheet: str, books: Sequence[str] = ()) -> set[SheetRef]:
    """First column of every cell/range reference in ``formula``."""
    tokens = _tokens(formula)
    if tokens is not None:
        candidates = [t.value for t in tokens if _is_range(t)]
    else:
        candidates = _REF_TOKEN_RE.findall(formula)
    refs: set[SheetRef] = set()
    for token in candidates:
        target = _range_first_column(token, default_sheet, books)
        if target is not None:
            refs.add(target)
    return refs
