"""Generate the clean 3-table fixture workbook (customer / order / order_item).

Orders and Order Items carry same-file VLOOKUP/XLOOKUP and arithmetic columns
so the formula pass has something to find. openpyxl does not compute formula
results, so those columns have no cached values until Excel saves the file.

Run ``python -m fixtures.clean_three_table out.xlsx`` or import ``generate``.
"""

from __future__ import annotations

import datetime as dt
import random
import sys
from pathlib import Path

from openpyxl import Workbook

CUSTOMER_ROWS = 25
ORDER_ROWS = 60
ITEM_ROWS = 150


def generate(path: str | Path, *, seed: int = 42) -> Path:
    """Write the workbook to ``path`` and return it."""
    rng = random.Random(seed)
    wb = Workbook()

    wb.remove(wb.worksheets[0])

    ws = wb.create_sheet("Customers")
    ws.append(["Customer ID", "Name", "Email", "Signup Date", "Active", "Credit Limit"])
    for i in range(1, CUSTOMER_ROWS + 1):
        ws.append(
            [
                i,
                f"Customer {i}",
                f"customer{i}@example.com",
                dt.date(2023, 1, 1) + dt.timedelta(days=rng.randint(0, 400)),
                rng.random() > 0.2,
                round(rng.uniform(500, 5000), 2),
            ]
        )

    ws = wb.create_sheet("Orders")
    ws.append(["Order ID", "Customer ID", "Ordered At", "Status", "Note", "Customer Name"])
    statuses = ["new", "paid", "shipped", "cancelled"]
    for i in range(1, ORDER_ROWS + 1):
        row = i + 1
        ws.append(
            [
                1000 + i,
                rng.randint(1, CUSTOMER_ROWS),
                dt.datetime(2024, 1, 1, 9, 0) + dt.timedelta(hours=rng.randint(0, 5000)),
                rng.choice(statuses),
                None if rng.random() < 0.7 else f"note {i}",
                f"=VLOOKUP(B{row},Customers!$A$2:$F${CUSTOMER_ROWS + 1},2,FALSE)",
            ]
        )

    ws = wb.create_sheet("Order Items")
    ws.append(["Item ID", "Order ID", "SKU", "Qty", "Unit Price", "Line Total", "Order Status"])
    for i in range(1, ITEM_ROWS + 1):
        row = i + 1
        ws.append(
            [
                i,
                1000 + rng.randint(1, ORDER_ROWS),
                f"SKU-{rng.randint(100, 999)}",
                rng.randint(1, 10),
                round(rng.uniform(1, 200), 2),
                f"=D{row}*E{row}",
                f"=_xlfn.XLOOKUP(B{row},Orders!$A$2:$A${ORDER_ROWS + 1},"
                f"Orders!$D$2:$D${ORDER_ROWS + 1})",
            ]
        )

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    print(generate(sys.argv[1] if len(sys.argv) > 1 else "clean_three_table.xlsx"))
