"""Generate a set of workbooks that reference each other (Phase 3 fixture).

``customers.xlsx``  Customers sheet, plain data.
``products.xlsx``   Products sheet, plain data.
``orders.xlsx``     Orders + Order Items; formulas pull from the other two
                    workbooks through external links (``[1]Customers!…``),
                    aggregate with SUMIF/COUNTIFS, use INDIRECT/OFFSET and
                    include a bare copy reference.
``summary.xlsx``    Per-customer totals aggregated from ``orders.xlsx``.

openpyxl does not compute formula results, so formula columns have no cached
values until Excel saves the files.

Run ``python -m fixtures.cross_file out_dir`` or import ``generate``.
"""

from __future__ import annotations

import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.packaging.relationship import Relationship
from openpyxl.workbook.external_link.external import (
    ExternalBook,
    ExternalLink,
    ExternalSheetNames,
)
from openpyxl.worksheet.worksheet import Worksheet

CUSTOMER_ROWS = 20
PRODUCT_ROWS = 10
ORDER_ROWS = 40
ITEM_ROWS = 100

CUSTOMERS = "customers.xlsx"
PRODUCTS = "products.xlsx"
ORDERS = "orders.xlsx"
SUMMARY = "summary.xlsx"


def _link(target: str, *sheets: str) -> ExternalLink:
    link = ExternalLink()
    link.file_link = Relationship(
        Id="rId1", type="externalLinkPath", Target=target, TargetMode="External"
    )
    link.externalBook = ExternalBook(sheetNames=ExternalSheetNames(sheetName=list(sheets)))
    return link


def _book(title: str) -> tuple[Workbook, Worksheet]:
    wb = Workbook()
    wb.remove(wb.worksheets[0])
    return wb, wb.create_sheet(title)


def _add_links(wb: Workbook, *links: ExternalLink) -> None:
    # openpyxl writes xl/externalLinks/externalLinkN.xml for every entry.
    wb._external_links.extend(links)  # type: ignore[attr-defined]


def generate(out_dir: str | Path, *, seed: int = 7) -> list[Path]:
    """Write the four workbooks under ``out_dir`` and return their paths."""
    rng = random.Random(seed)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []

    wb, ws = _book("Customers")
    ws.append(["Customer ID", "Name", "Region"])
    for i in range(1, CUSTOMER_ROWS + 1):
        ws.append([i, f"Customer {i}", rng.choice(["north", "south", "east", "west"])])
    paths.append(out / CUSTOMERS)
    wb.save(paths[-1])

    wb, ws = _book("Products")
    ws.append(["Product ID", "Product Name", "Unit Price"])
    for i in range(1, PRODUCT_ROWS + 1):
        ws.append([f"P{i:03d}", f"Product {i}", round(rng.uniform(5, 200), 2)])
    paths.append(out / PRODUCTS)
    wb.save(paths[-1])

    wb, ws = _book("Orders")
    _add_links(wb, _link(CUSTOMERS, "Customers"), _link(PRODUCTS, "Products"))
    ws.append(
        [
            "Order ID",
            "Customer ID",
            "Status",
            "Customer Name",
            "Item Count",
            "Order Total",
            "Status Copy",
            "Dynamic Note",
        ]
    )
    for i in range(1, ORDER_ROWS + 1):
        r = i + 1
        ws.append(
            [
                1000 + i,
                rng.randint(1, CUSTOMER_ROWS),
                rng.choice(["new", "paid", "shipped"]),
                f"=VLOOKUP(B{r},[1]Customers!$A$2:$C${CUSTOMER_ROWS + 1},2,FALSE)",
                f"=COUNTIFS('Order Items'!$B$2:$B${ITEM_ROWS + 1},A{r})",
                f"=SUMIF('Order Items'!$B$2:$B${ITEM_ROWS + 1},A{r},"
                f"'Order Items'!$E$2:$E${ITEM_ROWS + 1})",
                f"=C{r}",
                '=INDIRECT("Orders!C"&ROW())',
            ]
        )
    ws = wb.create_sheet("Order Items")
    ws.append(["Item ID", "Order ID", "Product ID", "Qty", "Line Total", "Next Qty"])
    for i in range(1, ITEM_ROWS + 1):
        r = i + 1
        ws.append(
            [
                i,
                1000 + rng.randint(1, ORDER_ROWS),
                f"P{rng.randint(1, PRODUCT_ROWS):03d}",
                rng.randint(1, 9),
                f"=D{r}*_xlfn.XLOOKUP(C{r},[2]Products!$A$2:$A${PRODUCT_ROWS + 1},"
                f"[2]Products!$C$2:$C${PRODUCT_ROWS + 1})",
                f"=OFFSET(D{r},1,0)",
            ]
        )
    paths.append(out / ORDERS)
    wb.save(paths[-1])

    wb, ws = _book("Summary")
    _add_links(wb, _link(ORDERS, "Orders"))
    ws.append(["Customer ID", "Orders", "Revenue"])
    for i in range(1, CUSTOMER_ROWS + 1):
        r = i + 1
        ws.append(
            [
                i,
                f"=COUNTIF([1]Orders!$B$2:$B${ORDER_ROWS + 1},A{r})",
                f"=SUMIF([1]Orders!$B$2:$B${ORDER_ROWS + 1},A{r},"
                f"[1]Orders!$F$2:$F${ORDER_ROWS + 1})",
            ]
        )
    paths.append(out / SUMMARY)
    wb.save(paths[-1])
    return paths


if __name__ == "__main__":
    for p in generate(sys.argv[1] if len(sys.argv) > 1 else "cross_file"):
        print(p)
