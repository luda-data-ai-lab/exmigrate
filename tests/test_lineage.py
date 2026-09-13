"""Phase 3: external references, aggregate/dynamic formulas, Lineage IR, include/exclude."""

from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import pytest
from flask.testing import FlaskClient
from openpyxl import Workbook, load_workbook

from exmigrate.analyzer import analyze_with_data
from exmigrate.analyzer.formulas import (
    OP_AGGREGATE,
    OP_CALCULATE,
    OP_COPY,
    OP_JOIN,
    SheetRef,
    external_books,
    scan_formulas,
)
from exmigrate.cli import main
from exmigrate.contracts.ir import SchemaIR
from exmigrate.contracts.lineage import LineageEdge, LineageIR, LineageNode
from exmigrate.lineage import to_flowchart
from exmigrate.service import apply_exclusions


def _edges(lineage: LineageIR) -> set[tuple[str, str, str]]:
    return {(e.source, e.target, e.op) for e in lineage.edges}


def _by_file(paths: list[Path]) -> dict[str, Path]:
    return {p.name: p for p in paths}


# --- formula pass -----------------------------------------------------------


def test_external_books_and_refs(cross_file_workbooks: list[Path]) -> None:
    orders = _by_file(cross_file_workbooks)["orders.xlsx"]
    wb = load_workbook(orders, read_only=True)
    try:
        assert external_books(wb) == ["customers.xlsx", "products.xlsx"]
    finally:
        wb.close()

    facts = scan_formulas(orders, {"Orders": 1, "Order Items": 1})
    orders_facts = facts["Orders"]
    # customer_name (D, idx 3) = VLOOKUP(..., [1]Customers!...)
    lookup = next(e for e in orders_facts.lookups if e.lookup_column == 1)
    assert lookup.target == SheetRef("Customers", 0, book="customers.xlsx")
    assert lookup.function == "VLOOKUP"
    assert SheetRef("Customers", 0, "customers.xlsx").sheet_label == "[customers.xlsx]Customers"
    # item_count (E, idx 4) = COUNTIFS('Order Items'!B:B, A) → aggregate from Order Items col 1
    assert (SheetRef("Order Items", 1), OP_AGGREGATE) in orders_facts.flows[4]
    # status_copy (G, idx 6) = C2 → copy from Orders col 2
    assert orders_facts.flows[6] == {(SheetRef("Orders", 2), OP_COPY): 40}
    # dynamic_note (H, idx 7) = INDIRECT(...)
    assert orders_facts.dynamic[7] == 40
    items = facts["Order Items"]
    assert items.dynamic[5] == 100
    assert (SheetRef("Products", 2, "products.xlsx"), OP_JOIN) in items.flows[4]


def test_full_column_and_arithmetic_refs(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.create_sheet("Data")
    ws.append(["k", "v", "total", "share"])
    for i in range(1, 6):
        ws.append([i, i * 10, f"=SUMIF($A:$A,A{i + 1},$B:$B)", f"=B{i + 1}/SUM($B$2:$B$6)"])
    path = tmp_path / "agg.xlsx"
    wb.save(path)
    facts = scan_formulas(path, {"Data": 1})["Data"]
    assert facts.flows[2] == {
        (SheetRef("Data", 0), OP_AGGREGATE): 5,
        (SheetRef("Data", 1), OP_AGGREGATE): 5,
    }
    # B2 and $B$2:$B$6 are the same source column: counted once per formula
    assert facts.flows[3] == {(SheetRef("Data", 1), OP_CALCULATE): 5}
    assert facts.functions[2] == {"SUMIF": 5} and facts.lookups == []


# --- analyzer integration -----------------------------------------------------


def test_cross_file_keys_and_issues(cross_file_workbooks: list[Path]) -> None:
    result = analyze_with_data(cross_file_workbooks)
    ir = result.ir
    orders = ir.table("orders")
    fk = orders.column("customer_id").fk
    assert fk is not None and (fk.table, fk.column, fk.confidence) == (
        "customers",
        "customer_id",
        0.99,
    )
    items_fk = ir.table("order_items").column("product_id").fk
    assert items_fk is not None and items_fk.table == "products" and items_fk.confidence == 0.99
    assert orders.column("order_total").derived and orders.column("dynamic_note").derived

    codes = {i.code for i in result.issues}
    assert "dynamic_reference" in codes and "missing_referenced_workbook" not in codes

    inventory = result.formulas.by_file()["orders.xlsx"]
    name_col = next(c for c in inventory if c.column == "customer_name")
    assert name_col.references == ["[customers.xlsx]Customers"]
    assert result.formulas.function_totals("summary.xlsx") == {"COUNTIF": 20, "SUMIF": 20}


def test_missing_workbook_issue(cross_file_workbooks: list[Path]) -> None:
    orders = _by_file(cross_file_workbooks)["orders.xlsx"]
    result = analyze_with_data([orders])
    missing = [i for i in result.issues if i.code == "missing_referenced_workbook"]
    messages = {i.message for i in missing}
    assert any(m.startswith("missing referenced workbook: customers.xlsx") for m in messages)
    assert any(m.startswith("missing referenced workbook: products.xlsx") for m in messages)
    assert all("upload" in m for m in messages)
    assert all(i.severity.value == "warning" for i in missing)
    # no FK can be inferred to a workbook that was not uploaded
    assert result.ir.table("orders").column("customer_id").fk is None

    lineage = result.lineage
    missing_nodes = {n.id for n in lineage.nodes if n.missing}
    assert "column:customers.xlsx/Customers/B" in missing_nodes
    assert any(n.kind == "dynamic_reference" for n in lineage.nodes)
    file_level = lineage.at_level("file")
    labels = {n.id: n.label for n in file_level.nodes}
    assert labels["file:customers.xlsx"] == "customers.xlsx (missing)"
    assert labels["file:orders.xlsx"] == "orders.xlsx"


def test_lineage_levels(cross_file_workbooks: list[Path]) -> None:
    lineage = analyze_with_data(cross_file_workbooks).lineage
    assert lineage.level == "column"
    col_edges = _edges(lineage)
    assert (
        "column:customers.xlsx/Customers/name",
        "column:orders.xlsx/Orders/customer_name",
        OP_JOIN,
    ) in col_edges
    assert (
        "column:orders.xlsx/Order Items/line_total",
        "column:orders.xlsx/Orders/order_total",
        OP_AGGREGATE,
    ) in col_edges
    assert (
        "column:orders.xlsx/Orders/status",
        "column:orders.xlsx/Orders/status_copy",
        OP_COPY,
    ) in col_edges
    assert ("dynamic:orders.xlsx", "column:orders.xlsx/Orders/dynamic_note", "dynamic") in col_edges

    sheet = lineage.at_level("sheet")
    sheet_edges = {(e.source, e.target, e.op): e.formula_count for e in sheet.edges}
    # 40 VLOOKUPs × (key column + returned column)
    customers, orders = "sheet:customers.xlsx/Customers", "sheet:orders.xlsx/Orders"
    assert sheet_edges[(customers, orders, OP_JOIN)] == 80
    assert sheet_edges[(orders, "sheet:summary.xlsx/Summary", OP_AGGREGATE)] == 60
    # same-sheet copies collapse away at sheet level
    assert not any(s == t for s, t, _ in sheet_edges)

    file_level = lineage.at_level("file")
    assert {n.id for n in file_level.nodes} >= {
        "file:orders.xlsx",
        "file:customers.xlsx",
        "file:products.xlsx",
        "file:summary.xlsx",
        "dynamic:orders.xlsx",
    }
    assert ("file:products.xlsx", "file:orders.xlsx", OP_JOIN) in _edges(file_level)
    assert file_level.at_level("file").model_dump() == file_level.model_dump()

    chart = to_flowchart(file_level)
    assert chart.startswith("flowchart LR") and "INDIRECT / OFFSET" in chart
    assert '-->|"join ×' in chart and "-.->" in chart


def test_lineage_ir_roundtrip_uses_from_to_aliases() -> None:
    ir = LineageIR(
        nodes=[
            LineageNode(id="column:a.xlsx/S/x", kind="column", label="s.x", file="a.xlsx"),
            LineageNode(id="column:a.xlsx/S/y", kind="column", label="s.y", file="a.xlsx"),
        ],
        edges=[LineageEdge(source="column:a.xlsx/S/x", target="column:a.xlsx/S/y", op="copy")],
    )
    payload = json.loads(ir.model_dump_json(by_alias=True))
    assert payload["edges"][0] == {
        "from": "column:a.xlsx/S/x",
        "to": "column:a.xlsx/S/y",
        "op": "copy",
        "formula_count": 0,
    }
    assert LineageIR.model_validate(payload) == ir


# --- include / exclude --------------------------------------------------------


def test_for_migration_drops_excluded_columns(clean_workbook: Path) -> None:
    result = analyze_with_data([clean_workbook])
    ir = result.ir
    ir.table("orders").column("customer_name").include = False
    ir.table("orders").column("order_id").include = False  # referenced by order_items FK
    assert ir.excluded_columns() == [("orders", "order_id"), ("orders", "customer_name")]
    trimmed, data = apply_exclusions(ir, result.table_data())
    names = [c.name for c in trimmed.table("orders").columns]
    assert "customer_name" not in names and "order_id" not in names
    assert list(data["orders"].columns) == names
    # FK pointing at the dropped column is removed rather than left dangling
    assert trimmed.table("order_items").column("order_id").fk is None
    # original IR is untouched
    assert ir.table("orders").column("customer_name").include is False
    assert ir.table("order_items").column("order_id").fk is not None
    assert SchemaIR.model_validate(ir.model_dump()).excluded_columns() == ir.excluded_columns()


# --- web ------------------------------------------------------------------


def _upload(client: FlaskClient, workbooks: list[Path]) -> str:
    handles = [p.open("rb") for p in workbooks]
    try:
        res = client.post(
            "/api/upload",
            data={"files": [(fh, p.name) for fh, p in zip(handles, workbooks, strict=True)]},
            content_type="multipart/form-data",
        )
    finally:
        for fh in handles:
            fh.close()
    assert res.status_code == 201, res.get_json()
    job_id: str = res.get_json()["job_id"]
    return job_id


def test_lineage_api(client: FlaskClient, cross_file_workbooks: list[Path]) -> None:
    job_id = _upload(client, cross_file_workbooks)
    assert b"Data flow" in client.get(f"/jobs/{job_id}/review").data

    body = client.get(f"/api/jobs/{job_id}/lineage").get_json()
    assert body["level"] == "column"
    assert {"from", "to", "op", "formula_count"} <= set(body["edges"][0])
    assert body["mermaid"].startswith("flowchart LR")

    for level, expected in (("file", "file"), ("sheet", "sheet"), ("col", "column")):
        res = client.get(f"/api/jobs/{job_id}/lineage?level={level}")
        assert res.status_code == 200 and res.get_json()["level"] == expected
    file_body = client.get(f"/api/jobs/{job_id}/lineage?level=file").get_json()
    assert "file:customers.xlsx" in {n["id"] for n in file_body["nodes"]}

    assert client.get(f"/api/jobs/{job_id}/lineage?level=galaxy").status_code == 400
    assert client.get("/api/jobs/nope/lineage").status_code == 404

    # jobs analyzed before lineage existed return an empty graph
    (client.application.extensions["job_store"].path(job_id) / "lineage.json").unlink()
    body = client.get(f"/api/jobs/{job_id}/lineage").get_json()
    assert body["nodes"] == [] and body["edges"] == []


def test_upload_keeps_unicode_filenames(
    client: FlaskClient, clean_workbook: Path, tmp_path: Path
) -> None:
    korean = tmp_path / "주문.xlsx"
    korean.write_bytes(clean_workbook.read_bytes())
    job_id = _upload(client, [korean])
    status = client.get(f"/api/jobs/{job_id}/status").get_json()
    assert status["files"] == ["주문.xlsx"]
    schema = client.get(f"/api/jobs/{job_id}/schema").get_json()
    assert {t["source_file"] for t in schema["tables"]} == {"주문.xlsx"}

    with clean_workbook.open("rb") as fh:
        res = client.post(
            "/api/upload",
            data={"files": (fh, "../../evil/../x.xlsx")},
            content_type="multipart/form-data",
        )
    assert res.status_code == 201
    uploads = client.application.extensions["job_store"].uploads_dir(res.get_json()["job_id"])
    assert [p.name for p in uploads.iterdir()] == ["x.xlsx"]


def test_exclude_via_web_migration(client: FlaskClient, clean_workbook: Path) -> None:
    job_id = _upload(client, [clean_workbook])
    schema = client.get(f"/api/jobs/{job_id}/schema").get_json()
    orders = next(t for t in schema["tables"] if t["name"] == "orders")
    name_col = next(c for c in orders["columns"] if c["name"] == "customer_name")
    assert name_col["derived"] is True and name_col["include"] is True
    name_col["include"] = False
    res = client.put(f"/api/jobs/{job_id}/schema", json=schema)
    assert res.status_code == 200
    saved = next(
        c
        for t in res.get_json()["tables"]
        if t["name"] == "orders"
        for c in t["columns"]
        if c["name"] == "customer_name"
    )
    assert saved["include"] is False

    res = client.post(f"/api/jobs/{job_id}/migrate", json={"targets": ["sqlite"]})
    status = res.get_json()
    assert res.status_code == 202 and status["state"] == "done"
    report = status["reports"][0]
    assert [i["code"] for i in report["issues"]] == ["column_excluded"]
    assert report["issues"][0]["column"] == "customer_name"

    db_path = client.application.extensions["job_store"].artifacts_dir(job_id) / "migration.db"
    con = sqlite3.connect(db_path)
    try:
        cols = [r[1] for r in con.execute("PRAGMA table_info(orders)")]
        assert "customer_name" not in cols and "customer_id" in cols
        assert con.execute("SELECT COUNT(*) FROM orders").fetchone()[0] == 60
    finally:
        con.close()


# --- CLI ------------------------------------------------------------------


def test_cli_lineage(cross_file_workbooks: list[Path], capsys: pytest.CaptureFixture[str]) -> None:
    files = [str(p) for p in cross_file_workbooks]
    assert main(["lineage", "--level", "file", *files]) == 0
    out = capsys.readouterr().out
    assert out.startswith("flowchart LR") and "summary.xlsx" in out

    assert main(["lineage", "--json", *files]) == 0
    body = json.loads(capsys.readouterr().out)
    assert body["level"] == "column" and body["edges"][0].keys() >= {"from", "to", "op"}


def test_cli_migrate_exclude(
    clean_workbook: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    out = tmp_path / "out"
    code = main(
        [
            "migrate",
            str(clean_workbook),
            "--target",
            "sqlite",
            "--out",
            str(out),
            "--exclude",
            "orders.customer_name",
            "--exclude",
            "order_items.line_total",
        ]
    )
    assert code == 0
    printed = capsys.readouterr().out
    assert "orders.customer_name was not migrated" in printed
    con = sqlite3.connect(out / "migration.db")
    try:
        assert "customer_name" not in [r[1] for r in con.execute("PRAGMA table_info(orders)")]
        assert "line_total" not in [r[1] for r in con.execute("PRAGMA table_info(order_items)")]
    finally:
        con.close()

    assert main(["migrate", str(clean_workbook), "--target", "sqlite", "--exclude", "nope.x"]) == 2
