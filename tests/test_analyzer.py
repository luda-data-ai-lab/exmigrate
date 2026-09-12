from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from exmigrate.analyzer import analyze, analyze_with_data, bind_data
from exmigrate.analyzer.formulas import SheetRef, _lookup_edge, _split_args, function_names
from exmigrate.analyzer.naming import dedupe, to_identifier
from exmigrate.contracts.ir import ColumnType


def test_clean_fixture_schema(clean_workbook: Path) -> None:
    ir, issues = analyze([clean_workbook])
    assert {i.code for i in issues} == {"formula_cache_empty"}
    assert [t.name for t in ir.tables] == ["customers", "orders", "order_items"]

    customers = ir.table("customers")
    assert customers.row_count == 25
    assert customers.source_sheet == "Customers"
    types = {c.name: c.type for c in customers.columns}
    assert types == {
        "customer_id": ColumnType.INTEGER,
        "name": ColumnType.TEXT,
        "email": ColumnType.TEXT,
        "signup_date": ColumnType.DATE,
        "active": ColumnType.BOOLEAN,
        "credit_limit": ColumnType.FLOAT,
    }
    assert customers.column("customer_id").nullable is False
    assert customers.column("email").max_length is not None

    orders = ir.table("orders")
    assert orders.row_count == 60
    assert orders.column("ordered_at").type is ColumnType.DATETIME
    note = orders.column("note")
    assert note.nullable is True and note.null_ratio > 0


def test_values_only_analysis_has_no_keys(clean_workbook: Path) -> None:
    result = analyze_with_data([clean_workbook], infer_keys=False)
    assert result.issues == []
    for table in result.ir.tables:
        for col in table.columns:
            assert col.pk is False and col.fk is None and col.derived is False


def test_clean_fixture_keys_and_derived(clean_workbook: Path) -> None:
    ir, _ = analyze([clean_workbook])
    pks = {t.name: [c.name for c in t.columns if c.pk] for t in ir.tables}
    assert pks == {
        "customers": ["customer_id"],
        "orders": ["order_id"],
        "order_items": ["item_id"],
    }
    fks = {
        (t.name, c.name): (c.fk.table, c.fk.column, c.fk.confidence)
        for t in ir.tables
        for c in t.columns
        if c.fk is not None
    }
    assert fks == {
        ("orders", "customer_id"): ("customers", "customer_id", 0.98),
        ("order_items", "order_id"): ("orders", "order_id", 0.98),
    }
    derived = {(t.name, c.name) for t in ir.tables for c in t.columns if c.derived}
    assert derived == {
        ("orders", "customer_name"),
        ("order_items", "line_total"),
        ("order_items", "order_status"),
    }


def test_fk_from_values_and_name_without_formulas(tmp_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.worksheets[0])
    ws = wb.create_sheet("Region")
    ws.append(["Region Code", "Label"])
    for code in ("KR", "JP", "US"):
        ws.append([code, f"label {code}"])
    ws = wb.create_sheet("Store")
    ws.append(["Store No", "Region Code", "Size"])
    for i, code in enumerate(["KR", "KR", "JP", "US", "JP", "KR"]):
        ws.append([i + 1, code, i % 3])
    path = tmp_path / "regions.xlsx"
    wb.save(path)

    ir, _ = analyze([path])
    assert ir.table("region").column("region_code").pk is True
    assert ir.table("store").column("store_no").pk is True
    fk = ir.table("store").column("region_code").fk
    assert fk is not None and (fk.table, fk.confidence) == ("region", 0.9)
    assert ir.table("store").column("size").fk is None


def test_function_names() -> None:
    assert function_names("=_xlfn.XLOOKUP(B2,Orders!$A$2:$A$61,Orders!$D$2:$D$61)") == ["XLOOKUP"]
    assert function_names('=SUM(A1:A3)+if(B1>0,ROUND(B1,2),"x(")') == ["SUM", "IF", "ROUND"]
    assert function_names("='My Sheet'!A2*Sheet1!B2") == []
    assert function_names("='O''Brien ROUND(2026)'!A2+SUM('My (x)'!A1:A2)") == ["SUM"]


def test_formula_inventory(clean_workbook: Path) -> None:
    inv = analyze_with_data([clean_workbook]).formulas
    assert [(c.sheet, c.column) for c in inv.columns] == [
        ("Orders", "customer_name"),
        ("Order Items", "line_total"),
        ("Order Items", "order_status"),
    ]
    name, total, status = inv.columns
    assert name.functions == {"VLOOKUP": 60} and name.references == ["Customers"]
    assert name.formula_cells == name.row_count == 60 and name.derived
    assert total.functions == {} and total.references == [] and total.sample == "=D2*E2"
    assert status.functions == {"XLOOKUP": 150} and status.references == ["Orders"]
    assert status.sample.startswith("=_xlfn.XLOOKUP(B2,")
    assert inv.function_totals() == {"XLOOKUP": 150, "VLOOKUP": 60}
    assert list(inv.by_file()) == [clean_workbook.name]
    assert analyze_with_data([clean_workbook], infer_keys=False).formulas.columns == []


def test_formula_helpers() -> None:
    assert _split_args("VLOOKUP(B2,'Order Items'!$A$2:$B$9,2,FALSE)", 8) == [
        "B2",
        "'Order Items'!$A$2:$B$9",
        "2",
        "FALSE",
    ]
    edge = _lookup_edge("VLOOKUP", ["B2", "'Order Items'!$A$2:$B$9", "2", "FALSE"], "Orders")
    assert edge is not None
    assert edge.lookup_column == 1 and edge.target == SheetRef("Order Items", 0)
    assert _lookup_edge("VLOOKUP", ["B2", "$A$2:$B$9", "2"], "Orders") is None


def test_ir_round_trips_json(clean_workbook: Path) -> None:
    ir, _ = analyze([clean_workbook])
    dumped = ir.model_dump_json()
    assert type(ir).model_validate_json(dumped) == ir


def test_header_detection_skips_title_rows(tmp_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.worksheets[0])

    ws = wb.create_sheet("Report")
    ws.append(["Monthly report"])
    ws.append([])
    ws.append(["id", "amount", "when"])
    ws.append([1, 2.5, "x"])
    ws.append([2, 3, "y"])
    path = tmp_path / "title.xlsx"
    wb.save(path)

    result = analyze_with_data([path])
    (table,) = result.ir.tables
    assert table.header_row == 3
    assert table.row_count == 2
    assert [c.name for c in table.columns] == ["id", "amount", "when"]
    assert table.column("amount").type is ColumnType.FLOAT


def test_blank_and_duplicate_headers(tmp_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.worksheets[0])

    ws = wb.create_sheet("Dup")
    ws.append(["Name", "Name", None, "1st col"])
    ws.append(["a", "b", 1, 2])
    path = tmp_path / "dup.xlsx"
    wb.save(path)

    ir, issues = analyze([path])
    (table,) = ir.tables
    assert [c.name for c in table.columns] == ["name", "name_2", "col_3", "_1st_col"]
    assert any(i.code == "blank_header" for i in issues)


def test_sheet_without_header_is_skipped(tmp_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.worksheets[0])

    ws = wb.create_sheet("Numbers")
    for i in range(5):
        ws.append([i, i * 2])
    path = tmp_path / "nohdr.xlsx"
    wb.save(path)
    ir, issues = analyze([path])
    assert ir.tables == []
    assert [i.code for i in issues] == ["no_header"]


def test_unsupported_file(tmp_path: Path) -> None:
    path = tmp_path / "data.csv"
    path.write_text("a,b\n1,2\n")
    ir, issues = analyze([path])
    assert ir.tables == []
    assert issues[0].code == "unsupported_file"


def test_bind_data_uses_positions(clean_workbook: Path) -> None:
    result = analyze_with_data([clean_workbook])
    result.ir.tables[0].name = "kunde"
    result.ir.tables[0].columns[0].name = "kid"
    data = bind_data(result.ir, result.frames)
    assert list(data["kunde"].columns)[0] == "kid"
    assert len(data["kunde"]) == 25


def test_naming_helpers() -> None:
    assert to_identifier("Order Items") == "order_items"
    assert to_identifier("고객 번호") == "고객_번호"
    assert to_identifier("2024 Sales") == "_2024_sales"
    assert to_identifier("   ", fallback="x") == "x"
    assert dedupe(["a", "a", "a_2", "b"]) == ["a", "a_2", "a_2_2", "b"]
