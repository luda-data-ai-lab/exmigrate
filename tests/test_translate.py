"""Phase 4: formula → SQL view / pandas translation."""

from __future__ import annotations

import runpy
import sqlite3
from pathlib import Path

import pytest
from flask.testing import FlaskClient
from openpyxl import Workbook

from exmigrate.analyzer import analyze_with_data
from exmigrate.cli import main
from exmigrate.contracts.translation import Translation
from exmigrate.service import run_migration
from exmigrate.translate import translate
from exmigrate.translate.parser import Binary, Call, Num, ParseError, Ref, Str, parse

PRODUCTS = [("P1", 10.0, "a"), ("P2", 20.0, "b"), ("P3", 30.0, "a")]
SALES = [(1, "P1", 2), (2, "P2", 1), (3, "P1", 3), (4, "P3", 4)]


@pytest.fixture(scope="module")
def shop_workbook(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Two sheets covering lookups, aggregates, arithmetic, copies and dynamic refs."""
    wb = Workbook()
    wb.remove(wb.worksheets[0])
    ws = wb.create_sheet("Products")
    ws.append(["product_id", "price", "category", "qty_sold", "n_sales"])
    for i, product in enumerate(PRODUCTS, start=2):
        ws.append(
            [
                *product,
                f"=SUMIF('Sales Log'!$B$2:$B$5,A{i},'Sales Log'!$C$2:$C$5)",
                f"=COUNTIFS('Sales Log'!$B$2:$B$5,A{i})",
            ]
        )
    ws = wb.create_sheet("Sales Log")
    ws.append(["sale_id", "product_id", "qty", "price", "total", "qty_copy", "size", "note"])
    for i, sale in enumerate(SALES, start=2):
        ws.append(
            [
                *sale,
                f"=VLOOKUP(B{i},Products!$A$2:$C$4,2,FALSE)",
                f"=C{i}*D{i}",
                f"=C{i}",
                f'=IF(C{i}>2,"big","small")',
                '=INDIRECT("A"&ROW())',
            ]
        )
    path = tmp_path_factory.mktemp("shop") / "shop.xlsx"
    wb.save(path)
    return path


def _translation(path: Path, *exclude: str) -> Translation:
    result = analyze_with_data([path])
    for spec in exclude:
        table, _, column = spec.rpartition(".")
        result.ir.table(table).column(column).include = False
    return translate(result.ir, result.formulas)


def _by_name(translation: Translation) -> dict[str, str]:
    return {f"{c.table}.{c.column}": c.status for c in translation.columns}


# --- parser -------------------------------------------------------------------


def test_parse_sheet_book_and_operators() -> None:
    ast = parse(
        "=D2*_xlfn.XLOOKUP(C2,[2]Products!$A$2:$A$11,'[2]My Sheet'!$C$2:$C$11)",
        ["a.xlsx", "b.xlsx"],
    )
    assert isinstance(ast, Binary) and ast.op == "*"
    assert isinstance(ast.left, Ref) and ast.left.text == "D2" and ast.left.is_cell
    call = ast.right
    assert isinstance(call, Call) and call.name == "XLOOKUP"
    rng = call.args[1]
    assert isinstance(rng, Ref) and (rng.book, rng.sheet) == ("b.xlsx", "Products")
    assert (rng.col1, rng.row1, rng.col2, rng.row2) == (0, 2, 0, 11) and not rng.is_cell
    quoted = call.args[2]
    assert isinstance(quoted, Ref) and quoted.sheet == "My Sheet"


def test_parse_literals_and_errors() -> None:
    ast = parse('=IF(A1>=3,"x",1.5%)')
    assert isinstance(ast, Call) and isinstance(ast.args[1], Str) and ast.args[1].value == "x"
    assert isinstance(ast.args[0], Binary) and ast.args[0].op == ">="
    assert isinstance(parse("=42"), Num)
    with pytest.raises(ParseError):
        parse("=SUM(A1:B2")
    with pytest.raises(ParseError):
        parse("={1,2,3}")


# --- translation --------------------------------------------------------------


def test_statuses_and_expressions(shop_workbook: Path) -> None:
    tr = _translation(shop_workbook, "sales_log.price", "sales_log.total")
    assert _by_name(tr) == {
        "products.qty_sold": "ok",
        "products.n_sales": "ok",
        "sales_log.price": "ok",
        "sales_log.total": "ok",
        "sales_log.qty_copy": "ok",
        "sales_log.size": "ok",
        "sales_log.note": "unsupported",
    }
    cols = {f"{c.table}.{c.column}": c for c in tr.columns}
    price = cols["sales_log.price"]
    assert not price.include
    assert price.sql["sqlite"] == (
        '(SELECT r1."price" FROM "products" r1 WHERE r1."product_id" = t."product_id" LIMIT 1)'
    )
    assert price.pandas == "_lookup(sales_log['product_id'], products, 'product_id', 'price')"
    # total reads the excluded derived column price: the lookup gets inlined
    total = cols["sales_log.total"]
    assert "sales_log.price" in total.depends_on
    assert 'SELECT r1."price" FROM "products"' in total.sql["sqlite"]
    assert total.sql["postgres"] == total.sql["sqlite"]
    # aggregates over another table use its recompute view only when needed
    qty_sold = cols["products.qty_sold"]
    assert qty_sold.sql["sqlite"] == (
        '(SELECT COALESCE(SUM(r1."qty"), 0) FROM "sales_log" r1 '
        'WHERE r1."product_id" = t."product_id")'
    )
    assert cols["products.n_sales"].sql["postgres"].startswith("(SELECT COALESCE(COUNT(*), 0)")
    assert cols["sales_log.qty_copy"].sql["sqlite"] == 't."qty"'
    assert (
        cols["sales_log.size"].sql["sqlite"]
        == "CASE WHEN (t.\"qty\" > 2) THEN 'big' ELSE 'small' END"
    )
    note = cols["sales_log.note"]
    assert note.sql == {"sqlite": "NULL", "postgres": "NULL"} and note.pandas == "np.nan"
    assert note.formula == '=INDIRECT("A"&ROW())'
    assert any(n.startswith("TODO INDIRECT()") for n in note.notes)
    assert tr.counts() == {"ok": 6, "partial": 0, "unsupported": 1}


def test_views_recompute_excluded_and_twin_included(shop_workbook: Path, tmp_path: Path) -> None:
    result = analyze_with_data([shop_workbook])
    for column in ("price", "total"):
        result.ir.table("sales_log").column(column).include = False
    tr = translate(result.ir, result.formulas)
    views = tr.views["sqlite"]
    assert views.count("CREATE VIEW") == 2
    assert "-- total: =C2*D2" in views
    assert "-- TODO INDIRECT()" in views

    reports = run_migration(result.ir, result.table_data(), ["sqlite"], {}, tmp_path, tr)
    (report,) = reports
    assert report.ok and not [i for i in report.issues if i.code == "view_not_created"]
    assert [Path(a).name for a in report.artifacts] == [
        "migration.db",
        "recompute_sqlite.sql",
        "recompute.py",
    ]
    con = sqlite3.connect(tmp_path / "migration.db")
    try:
        assert "price" not in [r[1] for r in con.execute("PRAGMA table_info(sales_log)")]
        rows = con.execute(
            "SELECT sale_id, price, total, qty_copy_calc, size_calc, note_calc "
            "FROM sales_log_v ORDER BY 1"
        ).fetchall()
        assert rows == [
            (1, 10.0, 20.0, 2, "small", None),
            (2, 20.0, 20.0, 1, "small", None),
            (3, 10.0, 30.0, 3, "big", None),
            (4, 30.0, 120.0, 4, "big", None),
        ]
        rows = con.execute(
            "SELECT product_id, qty_sold_calc, n_sales_calc FROM products_v ORDER BY 1"
        ).fetchall()
        assert rows == [("P1", 5, 2), ("P2", 1, 1), ("P3", 4, 1)]
    finally:
        con.close()

    # the generated pandas script agrees with the views
    module = runpy.run_path(str(tmp_path / "recompute.py"))
    con = sqlite3.connect(tmp_path / "migration.db")
    try:
        tables = module["recompute"](module["load"](con))
    finally:
        con.close()
    sales = tables["sales_log"].sort_values("sale_id")
    assert sales["price"].tolist() == [10.0, 20.0, 10.0, 30.0]
    assert sales["total"].tolist() == [20.0, 20.0, 30.0, 120.0]
    assert sales["size_calc"].tolist() == ["small", "small", "big", "big"]
    assert sales["size"].isna().all()  # cached Excel value kept as-is
    assert sales["note_calc"].isna().all()
    products = tables["products"].sort_values("product_id")
    assert products["qty_sold_calc"].tolist() == [5, 1, 4]
    assert products["n_sales_calc"].tolist() == [2, 1, 1]


def test_postgres_dump_appends_views(shop_workbook: Path, tmp_path: Path) -> None:
    result = analyze_with_data([shop_workbook])
    tr = translate(result.ir, result.formulas)
    configs = {"postgres": {"mode": "dump"}}
    (report,) = run_migration(result.ir, result.table_data(), ["postgres"], configs, tmp_path, tr)
    assert report.ok
    dump = (tmp_path / "migration.sql").read_text(encoding="utf-8")
    assert dump.index("COMMIT") < dump.index('CREATE VIEW "sales_log_v"')
    assert (tmp_path / "recompute_postgres.sql").read_text(encoding="utf-8") == tr.views["postgres"]


def test_cross_file_lookups_and_dependency_order(cross_file_workbooks: list[Path]) -> None:
    result = analyze_with_data(cross_file_workbooks)
    result.ir.table("order_items").column("line_total").include = False
    tr = translate(result.ir, result.formulas)
    cols = {f"{c.table}.{c.column}": c for c in tr.columns}
    assert cols["orders.customer_name"].status == "ok"
    assert 'FROM "customers" r1' in cols["orders.customer_name"].sql["sqlite"]
    line_total = cols["order_items.line_total"]
    assert line_total.status == "ok" and 'FROM "products" r1' in line_total.sql["sqlite"]
    # order_total sums the excluded line_total → reads the recompute view
    order_total = cols["orders.order_total"]
    assert 'FROM "order_items_v" r1' in order_total.sql["sqlite"]
    assert "order_items.line_total" in order_total.depends_on
    assert cols["order_items.next_qty"].status == "unsupported"
    views = tr.views["postgres"]
    assert views.index('CREATE VIEW "order_items_v"') < views.index('CREATE VIEW "orders_v"')
    script = tr.script
    assert script.index("order_items['line_total'] =") < script.index(
        "orders['order_total_calc'] ="
    )


def test_translation_api(client: FlaskClient, shop_workbook: Path) -> None:
    with shop_workbook.open("rb") as fh:
        res = client.post(
            "/api/upload",
            data={"files": [(fh, shop_workbook.name)]},
            content_type="multipart/form-data",
        )
    assert res.status_code == 201
    job_id = res.get_json()["job_id"]
    res = client.get(f"/api/jobs/{job_id}/translation")
    assert res.status_code == 200
    payload = res.get_json()
    assert payload["counts"] == {"ok": 6, "partial": 0, "unsupported": 1}
    assert {c["column"] for c in payload["columns"]} >= {"price", "total", "note"}
    assert all(c["include"] for c in payload["columns"])
    res = client.get(f"/api/jobs/{job_id}/translation?format=sqlite")
    assert res.status_code == 200 and b'CREATE VIEW "sales_log_v"' in res.data
    res = client.get(f"/api/jobs/{job_id}/translation?format=pandas")
    assert res.status_code == 200 and b"def recompute(" in res.data
    assert client.get(f"/api/jobs/{job_id}/translation?format=nope").status_code == 400
    assert client.get("/api/jobs/missing/translation").status_code == 404

    res = client.post(f"/api/jobs/{job_id}/migrate", json={"targets": ["sqlite"]})
    assert res.status_code == 202
    (report,) = res.get_json()["reports"]
    assert report["artifacts"] == ["migration.db", "recompute_sqlite.sql", "recompute.py"]
    assert client.get(f"/api/jobs/{job_id}/artifacts/recompute.py").status_code == 200


def test_cli_translate(shop_workbook: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["translate", str(shop_workbook), "--exclude", "sales_log.price"]) == 0
    out = capsys.readouterr().out
    assert "== sales_log.price [ok, excluded]  =VLOOKUP(B2,Products!$A$2:$C$4,2,FALSE)" in out
    assert "== sales_log.note [unsupported, included]" in out
    assert "7 formula column(s): 6 ok, 0 partial, 1 unsupported" in out

    assert main(["translate", str(shop_workbook), "--format", "postgres"]) == 0
    assert 'CREATE VIEW "products_v"' in capsys.readouterr().out
    assert main(["translate", str(shop_workbook), "--format", "json"]) == 0
    assert Translation.model_validate_json(capsys.readouterr().out).counts()["ok"] == 6
    assert main(["translate", str(shop_workbook), "--exclude", "nope.x"]) == 2


def test_cli_migrate_no_views(shop_workbook: Path, tmp_path: Path) -> None:
    out = tmp_path / "out"
    code = main(
        ["migrate", str(shop_workbook), "--target", "sqlite", "--out", str(out), "--no-views"]
    )
    assert code == 0
    assert not (out / "recompute.py").exists()
    con = sqlite3.connect(out / "migration.db")
    try:
        names = [r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='view'")]
    finally:
        con.close()
    assert names == []
